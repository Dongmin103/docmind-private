"""Actual retained extraction and ordinary model handoff with synthetic documents."""
import ast
import asyncio
import io
import logging
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from types import ModuleType, SimpleNamespace

import pytest
from api.db.services.file_service import FileService
from common.constants import LLMType

ROOT = Path(__file__).resolve().parents[5]


def _bytes(kind):
    output = io.BytesIO()
    if kind == 'pdf':
        from reportlab.pdfgen.canvas import Canvas
        canvas = Canvas(output)
        canvas.bookmarkPage('intro')
        canvas.addOutlineEntry('Synthetic outline', 'intro')
        canvas.drawString(72, 720, 'Synthetic attachment text')
        canvas.save()
    elif kind == 'docx':
        from docx import Document
        document = Document()
        document.add_paragraph('Synthetic attachment text')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Synthetic table'
        document.save(output)
    elif kind == 'pptx':
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = 'Synthetic attachment text'
        presentation.save(output)
    else:
        return b'Synthetic attachment text'
    return output.getvalue()


@pytest.fixture
def app_context(monkeypatch):
    apps = ModuleType('api.apps')
    apps.current_user = None
    monkeypatch.setitem(sys.modules, 'api.apps', apps)


@pytest.mark.parametrize('kind', ['txt', 'pdf', 'docx', 'pptx'])
def test_actual_attachment_extraction_preserves_content_envelope(kind, app_context):
    content = FileService.parse('synthetic.' + kind, _bytes(kind))
    assert 'File: synthetic.' + kind in content
    assert 'Content as following:' in content
    assert 'Synthetic attachment text' in content
    if kind == 'docx':
        assert 'Synthetic table' in content


def _answer_functions(namespace):
    source = (ROOT / 'api/db/services/dialog_service.py').read_text()
    names = {'async_chat_solo', 'get_files_content', 'split_file_attachments', 'convert_last_user_msg_to_multimodal', '_parse_data_uri_or_b64', '_normalize_text_from_content'}
    nodes = [n for n in ast.parse(source).body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name in names]
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(ROOT / 'api/db/services/dialog_service.py'), 'exec'), namespace)


def test_actual_extraction_reaches_ordinary_model_with_text_and_image(monkeypatch, app_context):
    from PIL import Image
    image = io.BytesIO()
    Image.new('RGB', (12, 12), color='white').save(image, format='PNG')
    blobs = {'doc': _bytes('pdf'), 'image': image.getvalue()}
    files = [{'id': 'doc', 'created_by': 'synthetic', 'name': 'synthetic.pdf', 'mime_type': 'application/pdf'}, {'id': 'image', 'created_by': 'synthetic', 'name': 'synthetic.png', 'mime_type': 'image/png'}]
    monkeypatch.setattr(FileService, 'get_blob', staticmethod(lambda _owner, location: blobs[location]))
    captured = []
    class Model:
        def __init__(self, *_args, **_kwargs):
            pass
        async def async_chat(self, system, history, gen_conf, **kwargs):
            captured.append((system, history, gen_conf, kwargs))
            return 'Synthetic answer'
    namespace = {'FileService': FileService, 'LLMType': LLMType, 'LLMBundle': Model, 'get_tenant_default_model_by_type': lambda *_args: {'model_type': 'chat', 'llm_factory': 'OpenAI-API-Compatible'}, 're': re, '_DATA_URI_RE': re.compile(r'^data:(?P<mime>[^;]+);base64,(?P<b64>[A-Za-z0-9+/=\s]+)$'), 'datetime': datetime, 'timezone': timezone, 'time': time, 'logging': logging}
    _answer_functions(namespace)
    dialog = SimpleNamespace(llm_id='', tenant_llm_id=None, tenant_id='synthetic', prompt_config={'system': 'ordinary system', 'tts': True, 'tavily_api_key': 'retired'}, llm_setting={})
    async def complete():
        return [event async for event in namespace['async_chat_solo'](dialog, [{'role': 'user', 'content': 'question', 'files': files}], stream=False)]
    events = asyncio.run(complete())
    assert events[0]['answer'] == 'Synthetic answer'
    assert 'audio_binary' not in events[0]
    blocks = captured[0][1][-1]['content']
    assert 'Synthetic attachment text' in blocks[0]['text']
    assert blocks[1]['image_url']['url'].startswith('data:image/png;base64,')


def test_audio_upload_rejected_before_storage_or_task_creation(monkeypatch, app_context):
    writes = []
    monkeypatch.setattr(FileService, 'get_root_folder', lambda *_a: writes.append('folder'))
    with pytest.raises(ValueError, match='audio'):
        FileService.upload_document.__wrapped__(FileService, SimpleNamespace(parser_id='naive', parser_config={}), [SimpleNamespace(filename='synthetic.mp3')], 'synthetic')
    assert writes == []
